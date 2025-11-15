#!/usr/bin/env python3
"""
Input/Output一覧表生成器
各テストケースの入出力変数をExcelで管理する
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from typing import List, Dict, Set
import re
import subprocess
from datetime import datetime
from mcdc_analyzer import TruthTable, TestCase

# バージョン情報
VERSION = "1.0.0"

def get_git_revision():
    """Gitリビジョン情報を取得する"""
    try:
        result = subprocess.run(
            ['git', 'rev-parse', '--short', 'HEAD'],
            capture_output=True,
            text=True,
            timeout=5
        )
        if result.returncode == 0:
            return result.stdout.strip()
        else:
            return "unknown"
    except (subprocess.TimeoutExpired, FileNotFoundError, Exception):
        return "unknown"


class IOTableExcelGenerator:
    """Input/Output一覧表生成器クラス"""

    def __init__(self):
        self.workbook = None

    def _extract_variables(self, test_cases: List[TestCase]) -> Set[str]:
        """
        テストケースから変数を抽出する

        Args:
            test_cases: TestCaseのリスト

        Returns:
            変数名のセット
        """
        variables = set()

        for test_case in test_cases:
            # 条件式から変数を抽出
            # パターン: 変数名（英数字とアンダースコア）
            matches = re.findall(r'\b([a-zA-Z_][a-zA-Z0-9_]*)\b', test_case.condition_expression)

            # 予約語と演算子を除外
            keywords = {'if', 'else', 'while', 'for', 'switch', 'case', 'default',
                       'return', 'sizeof', 'true', 'false', 'NULL'}

            for match in matches:
                if match not in keywords:
                    variables.add(match)

        return variables

    def create_io_table_excel(self, truth_table: TruthTable, output_path: str,
                              input_vars: List[str] = None, output_vars: List[str] = None):
        """
        Input/Output一覧表をExcelファイルとして生成する

        Args:
            truth_table: TruthTable オブジェクト
            output_path: 出力ファイルパス
            input_vars: 入力変数のリスト（Noneの場合は自動抽出）
            output_vars: 出力変数のリスト（Noneの場合は空）
        """
        self.workbook = Workbook()
        ws = self.workbook.active
        ws.title = "IO Table"

        # リビジョン情報を取得
        revision = get_git_revision()
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # 変数を抽出
        if input_vars is None:
            all_vars = sorted(self._extract_variables(truth_table.test_cases))
            # 簡易的に全てを入力変数として扱う
            input_vars = all_vars
            output_vars = []
        else:
            if output_vars is None:
                output_vars = []

        # スタイル定義
        header_font = Font(bold=True, size=11)
        info_font = Font(size=9, italic=True)
        header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        input_fill = PatternFill(start_color="E0F0FF", end_color="E0F0FF", fill_type="solid")
        output_fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # メタ情報行（1行目）
        total_cols = 2 + len(input_vars) + len(output_vars)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 4))
        meta_cell = ws['A1']
        meta_cell.value = f"C Unit Test Generator v{VERSION} | Revision: {revision} | Generated: {timestamp}"
        meta_cell.font = info_font
        meta_cell.alignment = center_alignment

        # ヘッダー行1: INPUT/OUTPUT（2行目に移動）
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
        cell = ws.cell(row=2, column=1)
        cell.value = ""
        cell.border = border

        # INPUT列のマージ（2行目）
        input_start_col = 3
        input_end_col = input_start_col + len(input_vars) - 1 if input_vars else input_start_col

        if input_vars:
            ws.merge_cells(start_row=2, start_column=input_start_col, end_row=2, end_column=input_end_col)
            cell = ws.cell(row=2, column=input_start_col)
            cell.value = "INPUT"
            cell.font = header_font
            cell.fill = input_fill
            cell.border = border
            cell.alignment = center_alignment

        # OUTPUT列のマージ（2行目）
        output_start_col = input_end_col + 1
        output_end_col = output_start_col + len(output_vars) - 1 if output_vars else output_start_col

        if output_vars:
            ws.merge_cells(start_row=2, start_column=output_start_col, end_row=2, end_column=output_end_col)
            cell = ws.cell(row=2, column=output_start_col)
            cell.value = "OUTPUT"
            cell.font = header_font
            cell.fill = output_fill
            cell.border = border
            cell.alignment = center_alignment

        # ヘッダー行3: 変数名（3行目に移動）
        headers = ["No", "テスト名"] + input_vars + output_vars

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header
            cell.font = header_font

            if col_num <= 2:
                cell.fill = header_fill
            elif col_num <= 2 + len(input_vars):
                cell.fill = input_fill
            else:
                cell.fill = output_fill

            cell.border = border
            cell.alignment = center_alignment

        # 列幅を設定
        ws.column_dimensions['A'].width = 6   # No
        ws.column_dimensions['B'].width = 30  # テスト名

        for i, var in enumerate(input_vars, 3):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = 12

        for i, var in enumerate(output_vars, 3 + len(input_vars)):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = 12

        # データ行（4行目から）
        for row_num, test_case in enumerate(truth_table.test_cases, 4):
            # No
            cell = ws.cell(row=row_num, column=1)
            cell.value = test_case.test_number
            cell.border = border
            cell.alignment = center_alignment

            # テスト名を生成
            condition_name = self._sanitize_test_name(test_case.condition_expression)
            test_name = f"test_{test_case.test_number:02d}_{condition_name}_{test_case.truth_values}"

            cell = ws.cell(row=row_num, column=2)
            cell.value = test_name
            cell.border = border
            cell.alignment = left_alignment

            # INPUT列（"-"で初期化）
            for i, var in enumerate(input_vars, 3):
                cell = ws.cell(row=row_num, column=i)
                cell.value = "-"
                cell.border = border
                cell.alignment = center_alignment

            # OUTPUT列（"-"で初期化）
            for i, var in enumerate(output_vars, 3 + len(input_vars)):
                cell = ws.cell(row=row_num, column=i)
                cell.value = "-"
                cell.border = border
                cell.alignment = center_alignment

        # 行の高さ
        ws.row_dimensions[1].height = 20  # メタ情報行
        ws.row_dimensions[2].height = 20  # INPUT/OUTPUT行
        ws.row_dimensions[3].height = 25  # 変数名行

        for row in range(4, len(truth_table.test_cases) + 4):
            ws.row_dimensions[row].height = 25

        # 保存
        self.workbook.save(output_path)
        print(f"I/O table saved to: {output_path}")

    def _sanitize_test_name(self, condition: str) -> str:
        """
        条件式からテスト名用の文字列を生成する
        （unity_test_generatorと同じロジック）
        """
        name = condition
        name = re.sub(r'==', '_eq_', name)
        name = re.sub(r'!=', '_ne_', name)
        name = re.sub(r'<=', '_le_', name)
        name = re.sub(r'>=', '_ge_', name)
        name = re.sub(r'<', '_lt_', name)
        name = re.sub(r'>', '_gt_', name)
        name = re.sub(r'\|\|', '_or_', name)
        name = re.sub(r'&&', '_and_', name)
        name = re.sub(r'[^a-zA-Z0-9_]', '_', name)
        name = re.sub(r'_+', '_', name)
        name = name.strip('_')

        if len(name) > 40:
            name = name[:40]

        return name


def main():
    """テスト用メイン関数"""
    from mcdc_analyzer import TruthTable, TestCase

    # テスト用の真偽表を作成
    truth_table = TruthTable(function_name="Utf1")

    test_cases = [
        TestCase(1, "(f4() & 0xdf) != 0", "T", "v9が7"),
        TestCase(2, "(f4() & 0xdf) != 0", "F", "v9!=7"),
        TestCase(3, "(mx63 == m47) || (mx63 == m46)", "TF", ""),
        TestCase(4, "(mx63 == m47) || (mx63 == m46)", "FT", ""),
        TestCase(5, "(mx63 == m47) || (mx63 == m46)", "FF", ""),
    ]

    truth_table.test_cases = test_cases

    # 入力変数と出力変数を指定（例）
    input_vars = ["Utv10", "Utv14", "Utv24", "mx63", "m47", "m46"]
    output_vars = ["Utv30", "Utv31", "Utv32"]

    # Excelファイル生成
    generator = IOTableExcelGenerator()
    generator.create_io_table_excel(truth_table, "test_io_table.xlsx",
                                   input_vars=input_vars, output_vars=output_vars)
    print("Test I/O table created successfully!")


if __name__ == "__main__":
    main()
