#!/usr/bin/env python3
"""
MC/DC 単体テスト自動生成ツール (純粋Python版)
Claude API 不使用 - 純粋なアルゴリズム実装

使い方: python3 generator/generate.py <Cソースファイルパス>

出力:
  work/<func_name>_truth_table.xlsx  - MC/DC真偽表
  work/<func_name>_test.c            - Unity単体テストコード
"""

import itertools
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

# work/ ディレクトリはスクリプトの2つ上のディレクトリ
WORK_DIR = Path(__file__).parent.parent / "work"


# ─────────────────────────── C コードパーサー ───────────────────────────

def parse_function_signature(source: str):
    """関数名とパラメータ名を抽出する"""
    m = re.search(r'\bint\s+(\w+)\s*\(([^)]*)\)', source)
    if not m:
        return None, []
    func_name = m.group(1)
    params_raw = m.group(2)
    params = []
    for p in params_raw.split(','):
        p = p.strip()
        if p:
            # "int a" → "a"
            parts = p.split()
            params.append(parts[-1].strip('*'))
    return func_name, params


def parse_if_conditions(source: str):
    """
    if文の条件式を抽出する（else if は除く）
    ネストされたif文も含めて順番に返す
    """
    conditions = []
    # else if を除いた if 文を探す
    for m in re.finditer(r'(?<![a-z_])if\s*\(', source):
        start = m.end()
        # 対応する閉じ括弧を探す
        depth = 1
        i = start
        while i < len(source) and depth > 0:
            if source[i] == '(':
                depth += 1
            elif source[i] == ')':
                depth -= 1
            i += 1
        cond = source[start:i - 1].strip()
        # else if を除外
        pre = source[max(0, m.start() - 5):m.start()].strip()
        if pre.endswith('else'):
            continue
        conditions.append(cond)
    return conditions


def parse_atomic_conditions(cond_str: str):
    """
    条件式からアトミック条件（単一比較）を抽出する
    例: "a > 0 && b > 0" → [A=(a>0), B=(b>0)]
    """
    pattern = r'(\w+)\s*(>=|<=|!=|==|>|<)\s*(-?\d+)'
    conditions = []
    for i, m in enumerate(re.finditer(pattern, cond_str)):
        var = m.group(1)
        op = m.group(2)
        threshold = int(m.group(3))
        cond_id = chr(ord('A') + i)
        expr = f"{var} {op} {threshold}"
        conditions.append({
            'id': cond_id,
            'expression': expr,
            'variable': var,
            'operator': op,
            'threshold': threshold,
            'span': (m.start(), m.end()),
        })
    return conditions


def normalize_expression(cond_str: str, conditions: list) -> str:
    """
    条件式をアトミック条件のIDに置換する
    例: "a > 0 && b > 0" → "A && B"
    後ろから置換して位置ズレを防ぐ
    """
    result = cond_str
    # 後ろから置換してスパンのズレを防ぐ
    for cond in reversed(conditions):
        s, e = cond['span']
        result = result[:s] + cond['id'] + result[e:]
    return result.strip()


# ─────────────────────────── 真偽値評価 ───────────────────────────

def evaluate_expression(expr: str, cond_values: dict) -> bool:
    """
    正規化された論理式を評価する
    例: "A && B" with {"A": True, "B": False} → False
    """
    e = expr
    # 長い識別子から置換（部分一致を防ぐ）
    for cid in sorted(cond_values.keys(), key=len, reverse=True):
        e = re.sub(r'\b' + re.escape(cid) + r'\b', str(cond_values[cid]), e)
    # C演算子 → Python演算子
    e = e.replace('&&', ' and ').replace('||', ' or ')
    # "! " → "not " (否定演算子)
    e = re.sub(r'!\s*(?!=)', 'not ', e)
    try:
        return bool(eval(e))
    except Exception:
        return False


# ─────────────────────────── MC/DCアルゴリズム ───────────────────────────

def get_true_value(op: str, threshold: int) -> int:
    """条件を真にするC変数値を返す"""
    mapping = {
        '>':  threshold + 1,
        '<':  threshold - 1,
        '>=': threshold,
        '<=': threshold,
        '==': threshold,
        '!=': threshold + 1,
    }
    return mapping[op]


def get_false_value(op: str, threshold: int) -> int:
    """条件を偽にするC変数値を返す"""
    mapping = {
        '>':  threshold,
        '<':  threshold,
        '>=': threshold - 1,
        '<=': threshold + 1,
        '==': threshold + 1,
        '!=': threshold,
    }
    return mapping[op]


def compute_mcdc_test_cases(conditions: list, norm_expr: str):
    """
    MC/DCアルゴリズムでテストケースを生成する

    MC/DC要件: 各条件について、その条件だけが変化して決定結果も変化する
    テストペアが選択集合に存在する必要がある（ペアの両要素が必要）

    Returns:
        list of dict: テストケースリスト
    """
    n = len(conditions)
    all_combos = list(itertools.product([True, False], repeat=n))

    # 全組み合わせの評価結果
    combo_results = []
    for combo in all_combos:
        cond_values = {conditions[i]['id']: combo[i] for i in range(n)}
        decision = evaluate_expression(norm_expr, cond_values)
        combo_results.append((combo, decision))

    # 各条件の独立性ペアを探す
    independence_pairs = {}
    for ci, cond in enumerate(conditions):
        pairs = []
        for i, (combo_i, dec_i) in enumerate(combo_results):
            for j, (combo_j, dec_j) in enumerate(combo_results):
                if i >= j:
                    continue
                # 条件ci以外はすべて同じで、条件ciの値が異なり、決定結果が異なる
                only_ci_differs = all(
                    combo_i[k] == combo_j[k] for k in range(n) if k != ci
                )
                if only_ci_differs and combo_i[ci] != combo_j[ci] and dec_i != dec_j:
                    pairs.append((i, j))
        independence_pairs[cond['id']] = pairs

    # MC/DC: 各条件について少なくとも1つのペアの「両要素」を選択集合に含める
    selected_indices = set()
    for cid, pairs in independence_pairs.items():
        if not pairs:
            continue
        # 既存選択集合との重なりが最大のペアを選ぶ
        best_pair = None
        best_overlap = -1
        for (i, j) in pairs:
            # 既にペアが完全にカバーされているか確認
            if i in selected_indices and j in selected_indices:
                best_pair = None  # 既にカバー済み
                break
            overlap = (1 if i in selected_indices else 0) + \
                      (1 if j in selected_indices else 0)
            if overlap > best_overlap:
                best_overlap = overlap
                best_pair = (i, j)
        if best_pair is not None:
            selected_indices.add(best_pair[0])
            selected_indices.add(best_pair[1])

    # テストケースを構築
    test_cases = []
    for tc_num, combo_idx in enumerate(sorted(selected_indices), 1):
        combo, decision = combo_results[combo_idx]
        cond_values = {conditions[i]['id']: combo[i] for i in range(n)}

        # パラメータ値を決定
        param_values = {}
        for cond in conditions:
            is_true = cond_values[cond['id']]
            val = get_true_value(cond['operator'], cond['threshold']) if is_true \
                  else get_false_value(cond['operator'], cond['threshold'])
            param_values[cond['variable']] = val

        test_cases.append({
            'tc_id': f"TC{tc_num}",
            'condition_values': cond_values,
            'decision': decision,
            'param_values': param_values,
        })

    return test_cases


# ─────────────────────────── 関数シミュレーション ───────────────────────────

def _find_matching_paren(text: str, start: int) -> int:
    """start位置の'('に対応する')'の位置を返す"""
    depth = 1
    i = start + 1
    while i < len(text) and depth > 0:
        if text[i] == '(':
            depth += 1
        elif text[i] == ')':
            depth -= 1
        i += 1
    return i - 1  # 閉じ括弧の位置


def _find_matching_brace(text: str, start: int) -> int:
    """start位置の'{'に対応する'}'の位置を返す"""
    depth = 1
    i = start + 1
    while i < len(text) and depth > 0:
        if text[i] == '{':
            depth += 1
        elif text[i] == '}':
            depth -= 1
        i += 1
    return i - 1  # 閉じ括弧の位置


def simulate_function(source: str, param_values: dict) -> int:
    """
    Cのif文を順番に評価して、最初にマッチした return 値を返す
    マッチしなければ最後の return 値を返す
    括弧カウンティングでネストされた条件式にも対応する
    """
    # 関数本体を抽出（最初の { から最後の } まで）
    brace_start = source.find('{')
    if brace_start == -1:
        return 0
    brace_end = _find_matching_brace(source, brace_start)
    body = source[brace_start + 1:brace_end]

    # コメントを除去
    body = re.sub(r'//[^\n]*', '', body)
    body = re.sub(r'/\*[\s\S]*?\*/', '', body)

    # if文を括弧カウンティングで解析
    pos = 0
    while pos < len(body):
        # "if (" を探す（else if は除外）
        m_if = re.search(r'(?<![a-z_])if\s*\(', body[pos:])
        if not m_if:
            break

        abs_if_start = pos + m_if.start()
        abs_paren_open = pos + m_if.end() - 1  # '(' の位置

        # else if を除外
        pre = body[:abs_if_start].rstrip()
        if pre.endswith('else'):
            pos = abs_paren_open + 1
            continue

        # 条件式の閉じ括弧を括弧カウンティングで探す
        paren_close = _find_matching_paren(body, abs_paren_open)
        cond_str = body[abs_paren_open + 1:paren_close].strip()

        # その後の { ... } ブロックを探す
        rest = body[paren_close + 1:]
        m_brace = re.search(r'\s*\{', rest)
        if not m_brace:
            pos = paren_close + 1
            continue

        abs_brace_open = paren_close + 1 + rest.index('{')
        brace_close = _find_matching_brace(body, abs_brace_open)
        block = body[abs_brace_open + 1:brace_close]

        # 条件式を実際の値で評価
        cond_eval = cond_str
        for var, val in sorted(param_values.items(), key=lambda x: len(x[0]), reverse=True):
            cond_eval = re.sub(r'\b' + re.escape(var) + r'\b', str(val), cond_eval)
        cond_eval = cond_eval.replace('&&', ' and ').replace('||', ' or ')
        cond_eval = re.sub(r'!\s*(?!=)', 'not ', cond_eval)

        try:
            result = bool(eval(cond_eval))
        except Exception:
            result = False

        if result:
            ret_m = re.search(r'return\s+(-?\d+)\s*;', block)
            if ret_m:
                return int(ret_m.group(1))

        pos = brace_close + 1

    # デフォルトのreturn値（最後のreturn文）
    all_returns = re.findall(r'return\s+(-?\d+)\s*;', body)
    if all_returns:
        return int(all_returns[-1])
    return 0


# ─────────────────────────── Excel生成 ───────────────────────────

def generate_excel(func_name: str, parameters: list,
                   if_statements_data: list, output_path: Path):
    """真偽表をExcelファイルに保存する"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center = Alignment(horizontal="center")

    for stmt in if_statements_data:
        ws = wb.create_sheet(title=stmt['if_id'])
        conditions = stmt['conditions']
        test_cases = stmt['test_cases']

        headers = (
            ["TC"]
            + parameters
            + [c['expression'] for c in conditions]
            + ["Decision", "Return"]
        )

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col)
            ].width = max(12, len(str(h)) + 2)

        for row_idx, tc in enumerate(test_cases, 2):
            col = 1
            ws.cell(row=row_idx, column=col, value=tc['tc_id'])
            col += 1
            for p in parameters:
                ws.cell(row=row_idx, column=col, value=tc['param_values'].get(p, 0))
                col += 1
            for cond in conditions:
                val = tc['condition_values'].get(cond['id'], False)
                ws.cell(row=row_idx, column=col, value="T" if val else "F")
                col += 1
            ws.cell(row=row_idx, column=col, value="T" if tc['decision'] else "F")
            col += 1
            ws.cell(row=row_idx, column=col, value=tc.get('expected_return', ''))

    wb.save(output_path)


# ─────────────────────────── Unityテストコード生成 ───────────────────────────

def generate_unity_tests(func_name: str, parameters: list,
                         all_test_cases: list, output_path: Path):
    """Unityテストコードを生成する"""
    lines = [
        f'#include "unity.h"',
        f'#include "{func_name}.h"',
        '',
        'void setUp(void) {}',
        'void tearDown(void) {}',
        '',
    ]

    test_names = []
    for tc in all_test_cases:
        if_id = tc['if_id']
        tc_id = tc['tc_id']
        test_name = f"test_{func_name}_{if_id}_{tc_id}"
        test_names.append(test_name)

        param_str = ', '.join(str(tc['param_values'].get(p, 0)) for p in parameters)
        cond_expr = tc.get('condition_expression', '')

        lines.append(f'/* {if_id}: {cond_expr} */')
        lines.append(f'void {test_name}(void) {{')
        lines.append(f'    TEST_ASSERT_EQUAL_INT({tc["expected_return"]}, {func_name}({param_str}));')
        lines.append('}')
        lines.append('')

    lines.append('int main(void) {')
    lines.append('    UNITY_BEGIN();')
    for name in test_names:
        lines.append(f'    RUN_TEST({name});')
    lines.append('    return UNITY_END();')
    lines.append('}')

    output_path.write_text('\n'.join(lines) + '\n')


# ─────────────────────────── メイン処理 ───────────────────────────

def main():
    if len(sys.argv) < 2:
        print("使い方: python3 generate.py <Cソースファイルパス>")
        sys.exit(1)

    c_file = Path(sys.argv[1])
    if not c_file.exists():
        print(f"エラー: ファイルが存在しません: {c_file}")
        sys.exit(1)

    source = c_file.read_text()

    # 関数シグネチャ解析
    func_name, parameters = parse_function_signature(source)
    if not func_name:
        print("エラー: 関数シグネチャを解析できません")
        sys.exit(1)

    print(f"解析中: {func_name}({', '.join(parameters)})")

    # if文の条件式を抽出
    if_cond_strs = parse_if_conditions(source)
    if not if_cond_strs:
        print("エラー: if文が見つかりません")
        sys.exit(1)

    # 各if文を解析
    all_if_data = []
    all_test_cases_flat = []

    for i, cond_str in enumerate(if_cond_strs):
        if_id = f"IF{i + 1}"
        conditions = parse_atomic_conditions(cond_str)
        if not conditions:
            print(f"警告: {if_id} のアトミック条件を抽出できません: {cond_str}")
            continue

        norm_expr = normalize_expression(cond_str, conditions)
        print(f"  {if_id}: {cond_str} → 正規化: {norm_expr}")

        test_cases = compute_mcdc_test_cases(conditions, norm_expr)

        # 期待リターン値を関数シミュレーションで決定
        # 未設定のパラメータは 0 をデフォルトとして補完する（simulate_function の正確な評価のため）
        for tc in test_cases:
            full_params = {p: 0 for p in parameters}
            full_params.update(tc['param_values'])
            tc['param_values'] = full_params
            tc['expected_return'] = simulate_function(source, tc['param_values'])
            tc['if_id'] = if_id
            tc['condition_expression'] = cond_str

        all_if_data.append({
            'if_id': if_id,
            'conditions': conditions,
            'test_cases': test_cases,
            'condition_expression': cond_str,
        })
        all_test_cases_flat.extend(test_cases)

    # work/ ディレクトリを確保
    WORK_DIR.mkdir(exist_ok=True)

    # Excel出力
    excel_path = WORK_DIR / f"{func_name}_truth_table.xlsx"
    generate_excel(func_name, parameters, all_if_data, excel_path)
    print(f"生成: {excel_path}")

    # Unityテストコード出力
    c_path = WORK_DIR / f"{func_name}_test.c"
    generate_unity_tests(func_name, parameters, all_test_cases_flat, c_path)
    print(f"生成: {c_path}")


if __name__ == '__main__':
    main()
